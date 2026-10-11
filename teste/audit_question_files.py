"""Read 2026 question-file layouts stored in pluraal.arquivos."""

from collections import Counter
import io
import json
import os
import re

import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook


def main():
    load_dotenv('.env')
    pg = psycopg.connect(host=os.getenv('PGHOST'), port=os.getenv('PGPORT'),
                        dbname=os.getenv('PGDATABASE'), user=os.getenv('PGUSER'),
                        password=os.getenv('PGPASSWORD'))
    counts = Counter()
    headers = Counter()
    lengths = Counter()
    errors = []
    cur = pg.cursor(name='question_file_audit')
    cur.execute('''SELECT tipo_arquivo,nome_arquivo,arquivo
        FROM pluraal.arquivos
        WHERE ano='2026' AND tipo_arquivo IN
            ('MAPA DA AVALIACAO','POR QUESTAO')''')
    for typ, name, blob in cur:
        counts[typ] += 1
        try:
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            first = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
            header = tuple(str(x or '').strip() for x in first)
            headers[(typ, header)] += 1
            lengths[(typ, ws.max_row)] += 1
            wb.close()
        except Exception as exc:
            errors.append((typ, name, type(exc).__name__, str(exc)[:100]))
        total = sum(counts.values())
        if total % 1000 == 0:
            print('Lidos', total, flush=True)
    cur.close()
    pg.close()
    print('COUNTS', json.dumps(counts, ensure_ascii=True))
    print('HEADERS', json.dumps([(k[0], list(k[1]), v)
                                 for k, v in headers.most_common(30)], ensure_ascii=True))
    print('TOP_LENGTHS', json.dumps([(k[0], k[1], v)
                                    for k, v in lengths.most_common(20)], ensure_ascii=True))
    print('ERRORS', len(errors), json.dumps(errors[:20], ensure_ascii=True))


if __name__ == '__main__':
    main()
