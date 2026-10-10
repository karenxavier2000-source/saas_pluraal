"""Reconcile MySQL 2026 question/map facts with original PostgreSQL files.

The original XLSX payloads are stored in pluraal.arquivos, including those
whose names end in .csv. 2025 MySQL rows are left untouched.
"""

import argparse
from collections import Counter
from decimal import Decimal, InvalidOperation
import io
import os
import re
import unicodedata

import mysql.connector
import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook


def norm(value):
    clean = ''.join(ch for ch in unicodedata.normalize('NFKD', str(value or '').casefold())
                    if not unicodedata.combining(ch))
    return ' '.join(clean.replace('–', '-').replace('—', '-').split())


def informative(value):
    clean = str(value or '').strip()
    return None if clean in ('', '-', '—', '–') else clean


def as_int(value):
    if value in (None, '', '-', '—'):
        return None
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def accuracy(value):
    if value in (None, '', '-', '—'):
        return None
    try:
        n = Decimal(str(value).replace('%', '').replace(',', '.'))
    except InvalidOperation:
        return None
    return n / 100 if n > 1 else n


def source_catalog(pg, my):
    schools = {}
    with pg.cursor() as cur:
        cur.execute('SELECT DISTINCT escola FROM pluraal.aluno WHERE ano=%s', ('2026',))
        for (name,) in cur:
            schools[norm(name)] = name
        cur.execute('SELECT "ESCOLA" FROM pluraal.escolas_pluraal')
        for (name,) in cur:
            if name:
                schools.setdefault(norm(name), name)
        cur.execute('''SELECT id_simulado,id_questao,
                MAX(enunciado),MAX(habilidade),MAX(dificuldade),MAX(gabarito)
            FROM pluraal.questao_simulados WHERE ano='2026'
            GROUP BY 1,2''')
        questions = {(norm(exam), norm(question)): (text, skill, difficulty, answer)
                     for exam, question, text, skill, difficulty, answer in cur}
        cur.execute('''SELECT id_simulado,escola,
                CASE WHEN COUNT(DISTINCT serie)=1 THEN MAX(serie) END,
                CASE WHEN COUNT(DISTINCT nivel)=1 THEN MAX(nivel) END,
                CASE WHEN COUNT(DISTINCT nivel_serie)=1 THEN MAX(nivel_serie) END
            FROM pluraal.aluno WHERE ano='2026' GROUP BY 1,2''')
        series = {(norm(exam), norm(school)): (grade, level, level_grade)
                  for exam, school, grade, level, level_grade in cur}
    cur = my.cursor()
    try:
        cur.execute('SELECT ESCOLA FROM escolas_pluraal UNION SELECT DISTINCT escola FROM simulados')
        for (name,) in cur:
            if name:
                schools.setdefault(norm(name), name)
    finally:
        cur.close()
    return schools, questions, series


def file_identity(name, schools):
    codes = re.findall(r' - ([A-Z0-9]{8}) - ', name, re.I)
    if not codes:
        raise ValueError(f'Código de simulado ausente: {name}')
    exam = codes[-1]
    rest = norm(os.path.splitext(name.split(' - ' + exam + ' - ', 1)[-1])[0])
    candidates = [candidate for candidate in schools
                  if rest.startswith(candidate) or
                  (len(rest) >= 10 and candidate.startswith(rest))]
    if not candidates:
        raise ValueError(f'Escola não identificada: {name}')
    best = max(candidates, key=len)
    if len([x for x in candidates if len(x) == len(best)]) > 1:
        raise ValueError(f'Escola ambígua: {name}')
    return exam, schools[best]


def read_sources(pg, schools, questions, series):
    maps = {}
    details = {}
    duplicates = Counter()
    files = Counter()
    bad_rows = Counter()
    cur = pg.cursor(name='original_question_files')
    cur.execute('''SELECT tipo_arquivo,nome_arquivo,arquivo,
            data_da_ultima_insercao FROM pluraal.arquivos
        WHERE ano='2026' AND tipo_arquivo IN
            ('MAPA DA AVALIACAO','POR QUESTAO')
        ORDER BY data_da_ultima_insercao,id''')
    for typ, name, blob, timestamp in cur:
        files[typ] += 1
        exam, school = file_identity(name, schools)
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        for position, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=1):
            if typ == 'MAPA DA AVALIACAO':
                code = informative(row[0] if len(row) > 0 else None)
                subject = informative(row[1] if len(row) > 1 else None)
                skill = informative(row[2] if len(row) > 2 else None)
                difficulty = informative(row[5] if len(row) > 5 else None)
                answer = informative(row[6] if len(row) > 6 else None)
            else:
                code = informative(row[1] if len(row) > 1 else None)
                subject = informative(row[2] if len(row) > 2 else None)
                skill = informative(row[3] if len(row) > 3 else None)
                difficulty = informative(row[5] if len(row) > 5 else None)
                answer = informative(row[6] if len(row) > 6 else None)
            if not code:
                bad_rows[typ] += 1
                continue
            key = (norm(exam), norm(school), norm(code))
            global_question = questions.get((norm(exam), norm(code)), (None, None, None, None))
            label = f'QUESTAO {position}'
            if typ == 'MAPA DA AVALIACAO':
                value = (exam, school, 2026, code, label,
                         global_question[0], subject, skill or global_question[1],
                         difficulty or global_question[2], answer or global_question[3],
                         timestamp)
                if key in maps:
                    duplicates[typ] += 1
                maps[key] = value
            else:
                grade, level, level_grade = series.get((norm(exam), norm(school)),
                                                       (None, None, None))
                number = as_int(row[0] if len(row) > 0 else None)
                if number is not None:
                    label = f'QUESTAO {number}'
                result = accuracy(row[7] if len(row) > 7 else None)
                letters = [as_int(row[index] if len(row) > index else None)
                           for index in range(8, 13)]
                value = (exam, school, 2026, grade, level, level_grade,
                         code, label, global_question[0], subject,
                         skill or global_question[1], difficulty or global_question[2],
                         result, *letters, timestamp)
                if key in details:
                    duplicates[typ] += 1
                details[key] = value
        wb.close()
        count = sum(files.values())
        if count % 1000 == 0:
            print(f'Arquivos lidos: {count}', flush=True)
    cur.close()
    print('Arquivos:', dict(files), 'linhas sem código:', dict(bad_rows),
          'chaves repetidas:', dict(duplicates), flush=True)
    return maps, details


def target_keys(cur, table):
    cur.execute(f'''SELECT id,ID_SIMULADO,ESCOLA,ID_QUESTAO
        FROM `{table}` WHERE ANO=2026''')
    result = {}
    duplicates = 0
    for id_, exam, school, question in cur:
        key = (norm(exam), norm(school), norm(question))
        if key in result:
            duplicates += 1
        result.setdefault(key, []).append(id_)
    return result, duplicates


def sync_table(cur, table, source, apply):
    target, dup = target_keys(cur, table)
    source_keys = set(source)
    target_keys_set = set(target)
    missing = source_keys - target_keys_set
    stale = target_keys_set - source_keys
    print(f'{table}: origem={len(source_keys)}, destino={sum(map(len,target.values()))}, '
          f'faltam={len(missing)}, extras={sum(len(target[k]) for k in stale)}, '
          f'duplicações internas={dup}', flush=True)
    if not apply:
        print('  faltam por escola:', Counter(k[1] for k in missing).most_common(12), flush=True)
        print('  extras por escola:', Counter(k[1] for k in stale).most_common(12), flush=True)
        print('  exemplos faltantes:', list(sorted(missing))[:8], flush=True)
        print('  exemplos extras:', list(sorted(stale))[:8], flush=True)
        return
    if table == 'mapa_simulados':
        columns = ('ID_SIMULADO','ESCOLA','ANO','ID_QUESTAO','QUESTAO',
                   'ENUNCIADO','DISCIPLINA','HABILIDADE','DIFICULDADE',
                   'GABARITO','data_insercao')
    else:
        columns = ('ID_SIMULADO','ESCOLA','ANO','SERIE','NIVEL',
                   'NIVEL_SERIE','ID_QUESTAO','QUESTAO','ENUNCIADO',
                   'DISCIPLINA','HABILIDADE','DIFICULDADE','ACERTOS',
                   'A','B','C','D','E','data_insercao')
    placeholders = ','.join(['%s'] * len(columns))
    names = ','.join('`' + c + '`' for c in columns)
    assignments = ','.join('`' + c + '`=%s' for c in columns)
    insert_sql = f'INSERT INTO `{table}` ({names}) VALUES ({placeholders})'
    update_sql = f'UPDATE `{table}` SET {assignments} WHERE id=%s'
    inserts = []
    updates = []
    delete_ids = []
    for key, value in source.items():
        ids = target.get(key, [])
        if ids:
            updates.append((*value, ids[0]))
            delete_ids.extend(ids[1:])
        else:
            inserts.append(value)
    for key in stale:
        delete_ids.extend(target[key])
    # The original files are authoritative for 2026. Rows not represented
    # there, and duplicate rows sharing a natural key, are removed.
    for start in range(0, len(delete_ids), 500):
        batch = delete_ids[start:start + 500]
        cur.execute(f'DELETE FROM `{table}` WHERE id IN ({",".join(["%s"]*len(batch))})', batch)
    for start in range(0, len(updates), 1000):
        cur.executemany(update_sql, updates[start:start + 1000])
    for start in range(0, len(inserts), 1000):
        cur.executemany(insert_sql, inserts[start:start + 1000])
    cur.execute(f'SELECT count(*) FROM `{table}` WHERE ANO=2026')
    actual = cur.fetchone()[0]
    if actual != len(source):
        raise ValueError(f'Contagem pós-migração divergente em {table}: {actual}/{len(source)}')
    print(f'{table}: validado {actual} linhas de 2026', flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    load_dotenv('.env')
    pg = psycopg.connect(host=os.getenv('PGHOST'), port=os.getenv('PGPORT'),
                        dbname=os.getenv('PGDATABASE'), user=os.getenv('PGUSER'),
                        password=os.getenv('PGPASSWORD'))
    my = mysql.connector.connect(host=os.getenv('MYSQL_HOST'),
                                 port=int(os.getenv('MYSQL_PORT', '3306')),
                                 database=os.getenv('MYSQL_DB'),
                                 user=os.getenv('MYSQL_USER'),
                                 password=os.getenv('MYSQL_PASSWORD'), autocommit=False)
    try:
        schools, questions, series = source_catalog(pg, my)
        maps, details = read_sources(pg, schools, questions, series)
        cur = my.cursor()
        try:
            cur.execute('SELECT count(*) FROM mapa_simulados WHERE ANO=2025')
            map_2025 = cur.fetchone()[0]
            cur.execute('SELECT count(*) FROM questao_simulado WHERE ANO=2025')
            detail_2025 = cur.fetchone()[0]
            sync_table(cur, 'mapa_simulados', maps, args.apply)
            sync_table(cur, 'questao_simulado', details, args.apply)
            if args.apply:
                cur.execute('SELECT count(*) FROM mapa_simulados WHERE ANO=2025')
                if cur.fetchone()[0] != map_2025:
                    raise ValueError('2025 alterado em mapa_simulados')
                cur.execute('SELECT count(*) FROM questao_simulado WHERE ANO=2025')
                if cur.fetchone()[0] != detail_2025:
                    raise ValueError('2025 alterado em questao_simulado')
                my.commit()
                print('Arquivos originais de 2026 sincronizados; 2025 preservado.', flush=True)
            else:
                print('Prévia concluída; nenhum dado alterado.', flush=True)
        finally:
            cur.close()
    except Exception:
        my.rollback()
        raise
    finally:
        pg.close()
        my.close()


if __name__ == '__main__':
    main()
